#!/usr/bin/env python3
"""Rebuild blood_sugar_log.xlsx from data/readings.json.

Run from the repo root:  python3 scripts/build_xlsx.py

This mirrors the reading types and target ranges used by index.html (see the
TYPES object in that file) and produces the same three-sheet workbook layout
(Log / Targets / Summary) that a person can also fill in by hand.
"""
import json
import datetime
import zoneinfo
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
READINGS_PATH = REPO_ROOT / "data" / "readings.json"
OUT_PATH = REPO_ROOT / "blood_sugar_log.xlsx"
TZ = zoneinfo.ZoneInfo("Asia/Karachi")

# Must match the TYPES object in index.html.
TYPE_MAP = {
    "fasting": "Fasting",
    "post": "Post-meal",
    "random": "Random",
    "bedtime": "Bedtime",
}
TARGETS = [
    ("Fasting", 80, 130),
    ("Post-meal", 80, 180),
    ("Random", 80, 160),
    ("Bedtime", 90, 150),
]

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="1F4E78")
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="666666")
BASE_FONT = Font(name=FONT_NAME, size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

LAST_ROW = 501  # header + 500 rows, same as the hand-editable template


def load_readings():
    if not READINGS_PATH.exists():
        return []
    with open(READINGS_PATH) as f:
        raw = json.load(f)
    rows = []
    for r in raw:
        if not isinstance(r, dict):
            continue
        t = r.get("t")
        v = r.get("v")
        typ = TYPE_MAP.get(r.get("type"))
        if not isinstance(t, (int, float)) or not isinstance(v, (int, float)) or not typ:
            continue
        dt = datetime.datetime.fromtimestamp(t / 1000, tz=TZ)
        rows.append((dt.strftime("%Y-%m-%d %H:%M"), v, typ, r.get("note") or "", t))
    rows.sort(key=lambda row: row[4])
    return [(dt, v, typ, note) for dt, v, typ, note, _ in rows]


def status_formula(row):
    return (
        f'=IF(OR($B{row}="",$C{row}=""),"",'
        f'IF(AND($B{row}>=VLOOKUP($C{row},Targets!$A$2:$C$5,2,FALSE),'
        f'$B{row}<=VLOOKUP($C{row},Targets!$A$2:$C$5,3,FALSE)),"In range",'
        f'IF($B{row}<VLOOKUP($C{row},Targets!$A$2:$C$5,2,FALSE),"Low","High")))'
    )


def build():
    readings = load_readings()

    wb = openpyxl.Workbook()

    # ---------------- Targets sheet ----------------
    targets = wb.active
    targets.title = "Targets"
    targets_data = [("Type", "Low", "High")] + TARGETS
    for r, row in enumerate(targets_data, start=1):
        for c, val in enumerate(row, start=1):
            cell = targets.cell(row=r, column=c, value=val)
            cell.font = HEADER_FONT if r == 1 else BASE_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            if r == 1:
                cell.fill = HEADER_FILL
    targets.column_dimensions["A"].width = 14
    targets.column_dimensions["B"].width = 10
    targets.column_dimensions["C"].width = 10
    targets["A7"] = "Target ranges (mg/dL) match the reading types used on the Sugar Tracker web page."
    targets["A7"].font = NOTE_FONT

    # ---------------- Log sheet ----------------
    log = wb.create_sheet("Log")
    headers = ["Date/Time", "Glucose (mg/dL)", "Type", "Status", "Note"]
    widths = [20, 16, 14, 12, 40]
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        col = get_column_letter(c)
        log.column_dimensions[col].width = w
        cell = log.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    dv_type = DataValidation(type="list", formula1="=Targets!$A$2:$A$5", allow_blank=True)
    log.add_data_validation(dv_type)
    dv_type.add(f"C2:C{LAST_ROW}")

    for r in range(2, LAST_ROW + 1):
        for c in (1, 2, 5):
            cell = log.cell(row=r, column=c)
            cell.font = BASE_FONT
            cell.border = BORDER
            cell.fill = INPUT_FILL
        log.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm"

        type_cell = log.cell(row=r, column=3)
        type_cell.font = BASE_FONT
        type_cell.border = BORDER
        type_cell.fill = INPUT_FILL
        type_cell.alignment = Alignment(horizontal="center")

        status_cell = log.cell(row=r, column=4)
        status_cell.value = status_formula(r)
        status_cell.font = BASE_FONT
        status_cell.border = BORDER
        status_cell.alignment = Alignment(horizontal="center")

        idx = r - 2
        if idx < len(readings):
            dt, v, typ, note = readings[idx]
            log.cell(row=r, column=1, value=dt)
            log.cell(row=r, column=2, value=v)
            log.cell(row=r, column=3, value=typ)
            log.cell(row=r, column=5, value=note)

    green_fill = PatternFill("solid", fgColor="C6EFCE")
    green_font = Font(name=FONT_NAME, color="006100")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    red_font = Font(name=FONT_NAME, color="9C0006")
    yellow_fill = PatternFill("solid", fgColor="FFEB9C")
    yellow_font = Font(name=FONT_NAME, color="9C6500")

    status_range = f"D2:D{LAST_ROW}"
    log.conditional_formatting.add(
        status_range, CellIsRule(operator="equal", formula=['"In range"'], fill=green_fill, font=green_font)
    )
    log.conditional_formatting.add(
        status_range, CellIsRule(operator="equal", formula=['"Low"'], fill=red_fill, font=red_font)
    )
    log.conditional_formatting.add(
        status_range, CellIsRule(operator="equal", formula=['"High"'], fill=yellow_fill, font=yellow_font)
    )

    log.freeze_panes = "A2"

    # ---------------- Summary sheet ----------------
    summary = wb.create_sheet("Summary")
    summary["A1"] = "Blood Sugar Summary"
    summary["A1"].font = TITLE_FONT
    summary.merge_cells("A1:B1")

    rows = [
        ("Total Readings", f'=COUNTIF(Log!B2:B{LAST_ROW},"<>")'),
        ("Average Glucose (mg/dL)", f'=IFERROR(AVERAGE(Log!B2:B{LAST_ROW}),"")'),
        ("Estimated HbA1c (%)", '=IFERROR((B4+46.7)/28.7,"")'),
        ("% Readings In Range", f'=IFERROR(COUNTIF(Log!D2:D{LAST_ROW},"In range")/COUNTIF(Log!B2:B{LAST_ROW},"<>"),"")'),
        ("Readings Below 70 mg/dL", f'=COUNTIF(Log!B2:B{LAST_ROW},"<70")'),
    ]
    start_r = 3
    for i, (label, formula) in enumerate(rows):
        r = start_r + i
        lc = summary.cell(row=r, column=1, value=label)
        lc.font = BASE_FONT
        lc.border = BORDER
        vc = summary.cell(row=r, column=2, value=formula)
        vc.font = Font(name=FONT_NAME, size=11, bold=True)
        vc.border = BORDER
        vc.alignment = Alignment(horizontal="center")

    summary["B5"].number_format = "0.0"
    summary["B6"].number_format = "0.0%"
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 16

    summary["A9"] = "Auto-generated by scripts/build_xlsx.py from data/readings.json — do not hand-edit this file, edit the app or data/readings.json instead."
    summary["A9"].font = NOTE_FONT
    summary.merge_cells("A9:E9")

    wb._sheets = [wb["Log"], wb["Targets"], wb["Summary"]]
    wb.active = 0

    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} with {len(readings)} reading(s).")


if __name__ == "__main__":
    build()
